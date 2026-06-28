"""Excel parser for .xlsx (openpyxl) and .xls (pandas) workbooks.

Iterates every sheet. The header row is often not the first row (title/contact
rows precede it), so we scan the top rows with find_header_row before mapping
columns. Each sheet's data rows become ParsedRows tagged with source_ref=sheet.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from datetime import date, datetime

import openpyxl

from ..enums import FileFormat
from .base import BaseParser, ParsedDocument, register_parser
from .pdf_text import extract_header_hints
from .table_extract import rows_from_table


def _cellval(v: object) -> str:
    """Render an openpyxl/pandas cell value as a trimmed string."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float):
        if v != v:  # NaN — pandas renders blank .xls cells this way; treat as empty
            return ""
        if v.is_integer():
            return str(int(v))
    return str(v).strip()


def _level_from(first_col: int, leading_spaces: int, indent: float) -> int | None:
    """Combine the geometric indentation signals into a coarse nesting level.

    A header pushed further to the right is more deeply nested. We sum the column
    offset of its first non-empty cell, its leading-space count (~2 spaces/level)
    and the openpyxl cell-alignment indent. Returns None when there is no
    positive signal so table_extract falls back to keyword-class depth.
    """
    level = first_col + leading_spaces // 2 + int(indent)
    return level if level > 0 else None


def _cell_level_hint(cells) -> int | None:
    """Geometric nesting level for a row of openpyxl cells (or None).

    Uses the first non-empty cell's column index, the leading spaces of its raw
    text (NOT stripped, so indentation survives) and cell.alignment.indent.
    """
    for ci, cell in enumerate(cells):
        v = cell.value
        if v is None:
            continue
        raw = str(v)
        if not raw.strip():
            continue
        leading = len(raw) - len(raw.lstrip())
        indent = 0.0
        try:  # available even in read_only mode; guard for safety
            al = cell.alignment
            if al is not None and al.indent:
                indent = float(al.indent)
        except Exception:
            indent = 0.0
        return _level_from(ci, leading, indent)
    return None


def _value_level_hint(values) -> int | None:
    """Geometric nesting level for a row of raw values (legacy .xls / pandas)."""
    for ci, v in enumerate(values):
        if v is None:
            continue
        raw = str(v)
        if not raw.strip():
            continue
        leading = len(raw) - len(raw.lstrip())
        return _level_from(ci, leading, 0.0)
    return None


# --------------------------------------------------------------------------- #
# Legacy .xls -> .xlsx conversion via LibreOffice (Fix A fallback).            #
# --------------------------------------------------------------------------- #
def _find_soffice() -> str | None:
    """Locate the LibreOffice headless binary on PATH or in common install dirs."""
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for cand in (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
        "/opt/libreoffice/program/soffice",
        "/snap/bin/libreoffice",
    ):
        if os.path.exists(cand):
            return cand
    return None


def _convert_xls_to_xlsx(src_path: str) -> str | None:
    """Convert a legacy .xls to .xlsx via `soffice --headless`. Returns the new
    path (caller cleans up its temp dir) or None if conversion is unavailable."""
    soffice = _find_soffice()
    if soffice is None:
        return None
    outdir = tempfile.mkdtemp(prefix="medarchive_xls_")
    # Isolate the LibreOffice user profile so a headless run never collides with
    # a desktop instance or a read-only HOME.
    profile = f"file://{os.path.join(outdir, 'profile')}"
    cmd = [
        soffice, "--headless", f"-env:UserInstallation={profile}",
        "--convert-to", "xlsx", "--outdir", outdir, src_path,
    ]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=180)
    except Exception:
        shutil.rmtree(outdir, ignore_errors=True)
        return None
    out = os.path.join(outdir, os.path.splitext(os.path.basename(src_path))[0] + ".xlsx")
    if os.path.exists(out):
        return out
    shutil.rmtree(outdir, ignore_errors=True)
    return None


class _BaseExcelParser(BaseParser):
    def _finish(self, doc: ParsedDocument, top_text_parts: list[str]) -> ParsedDocument:
        text = "\n".join(p for p in top_text_parts if p)
        doc.raw_text = text
        extract_header_hints(text, doc)
        if not doc.rows:
            doc.add_warning("No priced rows extracted from workbook.")
        return doc


@register_parser
class XlsxParser(_BaseExcelParser):
    formats = (FileFormat.xlsx, FileFormat.xls)

    def parse(self, file_path: str, file_name: str | None = None) -> ParsedDocument:
        name = (file_name or file_path).lower()
        fmt = FileFormat.xls if name.endswith(".xls") else FileFormat.xlsx
        doc = ParsedDocument(file_format=fmt)
        if fmt is FileFormat.xls:
            return self._parse_xls(file_path, doc)
        return self._parse_xlsx(file_path, doc)

    # ---- .xlsx via openpyxl ------------------------------------------------ #
    def _parse_xlsx(self, file_path: str, doc: ParsedDocument) -> ParsedDocument:
        top_text: list[str] = []
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        except Exception as exc:
            doc.add_warning(f"openpyxl could not open workbook: {exc}")
            return doc
        try:
            for ws in wb.worksheets:
                # Iterate CELL objects (not values_only) so we can read both the
                # value and the alignment indent for the section-depth hint.
                table: list[list[str]] = []
                level_hints: list[int | None] = []
                for row in ws.iter_rows():
                    table.append([_cellval(c.value) for c in row])
                    level_hints.append(_cell_level_hint(row))
                if not table:
                    continue
                top_text.append(ws.title)
                for r in table[:6]:
                    line = " ".join(c for c in r if c)
                    if line:
                        top_text.append(line)
                doc.rows.extend(
                    rows_from_table(
                        table,
                        source_ref_prefix=f"sheet={ws.title}",
                        level_hints=level_hints,
                    )
                )
        finally:
            wb.close()
        return self._finish(doc, top_text)

    # ---- .xls via xlrd, with a LibreOffice conversion fallback ------------- #
    def _parse_xls(self, file_path: str, doc: ParsedDocument) -> ParsedDocument:
        sheets = self._read_xls_xlrd(file_path, doc)
        if sheets is None:
            # xlrd missing or failed -> convert .xls -> .xlsx and reuse the
            # openpyxl path (which preserves header detection / sections / tiers).
            converted = _convert_xls_to_xlsx(file_path)
            if converted is None:
                doc.add_warning(
                    "Cannot read legacy .xls: xlrd unavailable/failed and no "
                    "LibreOffice (soffice) found for fallback conversion. "
                    "Install xlrd>=2.0.1 or LibreOffice."
                )
                return doc
            doc.add_warning("Read .xls via LibreOffice .xls->.xlsx conversion fallback.")
            try:
                return self._parse_xlsx(converted, doc)
            finally:
                shutil.rmtree(os.path.dirname(converted), ignore_errors=True)

        top_text: list[str] = []
        for sheet_name, df in sheets.items():
            raw_rows = df.values.tolist()
            table = [[_cellval(v) for v in row] for row in raw_rows]
            if not table:
                continue
            level_hints = [_value_level_hint(row) for row in raw_rows]
            top_text.append(str(sheet_name))
            for r in table[:6]:
                line = " ".join(c for c in r if c)
                if line:
                    top_text.append(line)
            doc.rows.extend(
                rows_from_table(
                    table,
                    source_ref_prefix=f"sheet={sheet_name}",
                    level_hints=level_hints,
                )
            )
        return self._finish(doc, top_text)

    @staticmethod
    def _read_xls_xlrd(file_path: str, doc: ParsedDocument):
        """Read every sheet of a legacy .xls with the xlrd engine.

        Returns a ``{sheet_name: DataFrame}`` dict, or None to signal that the
        caller should fall back to LibreOffice conversion.
        """
        try:
            import pandas as pd  # local import: only needed for legacy .xls
        except Exception as exc:  # pragma: no cover
            doc.add_warning(f"pandas unavailable for .xls: {exc}")
            return None
        try:
            import xlrd  # noqa: F401  ensure the legacy engine is importable
        except Exception:
            doc.add_warning("xlrd not installed; trying LibreOffice fallback.")
            return None
        try:
            return pd.read_excel(
                file_path, sheet_name=None, header=None, dtype=str, engine="xlrd"
            )
        except Exception as exc:
            doc.add_warning(f"xlrd read failed ({exc}); trying LibreOffice fallback.")
            return None
